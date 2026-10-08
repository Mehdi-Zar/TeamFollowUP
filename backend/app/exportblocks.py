"""Studio des exports: the generic slides and widgets every document can carry.

A cover page, a free text slide, a free layout (a 12 x 8 grid whose zones hold
widgets), and a slide imported from a PowerPoint file with its ``{fields}`` filled
in. Everything here draws with python-pptx on a slide it is given, in inches, and
reads the active theme through ``pptxtpl``; the data-bound widgets (a squad's
timeline, a platform's KPI panel...) are drawn by the renderers that own them and
only placed here.
"""
from __future__ import annotations

import copy
import io

from . import pptxtpl
from .exportspec import GRID_COLS, GRID_ROWS, fill

SW, SH = pptxtpl.SLIDE_W_IN, pptxtpl.SLIDE_H_IN
MARGIN_X, MARGIN_TOP, MARGIN_BOTTOM = 0.4, 0.35, 0.35
TITLE_H = 0.95
ZONE_GAP = 0.14

_DEF = {"primary": "#1E2761", "ink": "#111827", "muted": "#55606E", "line": "#E2E8F0",
        "background": "#FFFFFF", "card": "#F1F5F9"}


def _c(key: str) -> str:
    return pptxtpl.color(key, _DEF[key])


def _rgb(hexstr: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexstr.lstrip("#").upper())


def _textbox(slide, x, y, w, h, paras, *, anchor="top", margin=0.0, fill_hex=None, wrap=True):
    """A text box; ``paras`` is a list of (text, size_pt, hex, bold, align)."""
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
    if fill_hex:
        from pptx.enum.shapes import MSO_SHAPE
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        box.fill.solid(); box.fill.fore_color.rgb = _rgb(fill_hex)
        box.line.fill.background(); box.shadow.inherit = False
    else:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                          "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin / 2)
    aligns = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    for i, (text, size, color, bold, align) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = aligns.get(align, PP_ALIGN.LEFT)
        p.space_after = Pt(4)
        r = p.add_run(); r.text = text
        r.font.size = Pt(pptxtpl.font_size(size)); r.font.bold = bool(bold)
        r.font.color.rgb = _rgb(color)
    return box


def _rect(slide, x, y, w, h, hexfill):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = _rgb(hexfill)
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh


def picture_in(slide, blob: bytes, x, y, w, h):
    """An image fitted inside the frame, its proportions kept, centred."""
    from pptx.util import Inches
    try:
        from PIL import Image  # noqa: F401  (python-pptx reads sizes itself)
    except Exception:  # pragma: no cover - Pillow ships with python-pptx
        pass
    pic = slide.shapes.add_picture(io.BytesIO(blob), Inches(x), Inches(y))
    iw, ih = pic.width or 1, pic.height or 1
    scale = min(Inches(w) / iw, Inches(h) / ih)
    pic.width, pic.height = int(iw * scale), int(ih * scale)
    pic.left = int(Inches(x) + (Inches(w) - pic.width) / 2)
    pic.top = int(Inches(y) + (Inches(h) - pic.height) / 2)
    return pic


def title_band(slide, title: str):
    """The navy band most decks open with, the title written in it."""
    _rect(slide, 0, 0, SW, TITLE_H, _c("primary"))
    _textbox(slide, 0.5, 0.0, SW - 1.0, TITLE_H, [(title, 24, "#FFFFFF", True, "left")], anchor="middle")


# --------------------------------------------------------------------------
# Generic sections
# --------------------------------------------------------------------------
def cover_slide(prs, params: dict, values: dict, image: bytes | None = None):
    """A page de garde: the background colour, a large title, a subtitle, a logo."""
    slide = pptxtpl.add_slide(prs)
    # A cover page carries no footer nor page number (see footer()).
    slide._element.cSld.set("name", COVER_NAME)
    bg = params.get("background") or _c("primary")
    _rect(slide, 0, 0, SW, SH, bg)
    if image:
        picture_in(slide, image, 0.6, 0.5, 3.0, 1.2)
    _textbox(slide, 0.9, 2.4, SW - 1.8, 1.6,
             [(fill(params.get("title") or "", values), 36, "#FFFFFF", True, "left")], anchor="bottom")
    sub = fill(params.get("subtitle") or "", values)
    if sub:
        _textbox(slide, 0.9, 4.15, SW - 1.8, 1.0, [(sub, 16, "#E5E7EB", False, "left")])
    return slide


def text_slide(prs, params: dict, values: dict):
    """A free text slide: a title band and paragraphs; a line starting with "- " is a
    bullet. Every ``{field}`` is filled in."""
    slide = pptxtpl.add_slide(prs)
    title = fill(params.get("title") or "", values)
    top = 0.6
    if title:
        title_band(slide, title)
        top = TITLE_H + 0.35
    paras = []
    for raw in fill(params.get("body") or "", values).splitlines():
        line = raw.rstrip()
        if not line:
            paras.append((" ", 8, _c("ink"), False, "left"))
        elif line.startswith("- "):
            paras.append(("• " + line[2:], 16, _c("ink"), False, "left"))
        elif line.startswith("# "):
            paras.append((line[2:], 20, _c("primary"), True, "left"))
        else:
            paras.append((line, 16, _c("ink"), False, "left"))
    if paras:
        _textbox(slide, 0.7, top, SW - 1.4, SH - top - 0.5, paras)
    return slide


def zone_rect(zone, has_title: bool) -> tuple[float, float, float, float]:
    """The frame (inches) of a grid zone [c0, r0, c1, r1], gaps taken out."""
    c0, r0, c1, r1 = zone
    top = (TITLE_H + 0.2) if has_title else MARGIN_TOP
    cw = (SW - 2 * MARGIN_X) / GRID_COLS
    rh = (SH - top - MARGIN_BOTTOM) / GRID_ROWS
    x = MARGIN_X + c0 * cw + ZONE_GAP / 2
    y = top + r0 * rh + ZONE_GAP / 2
    return x, y, (c1 - c0) * cw - ZONE_GAP, (r1 - r0) * rh - ZONE_GAP


def grid_slide(prs, sec: dict, values: dict, draw_widget) -> object:
    """A free layout: its title band, then each widget in its zone, in order.
    ``draw_widget(slide, item, x, y, w, h)`` draws one placed widget."""
    slide = pptxtpl.add_slide(prs)
    title = fill((sec.get("params") or {}).get("title") or "", values)
    if title:
        title_band(slide, title)
    for item in sec.get("items") or []:
        x, y, w, h = zone_rect(item["zone"], bool(title))
        draw_widget(slide, item, x, y, w, h)
    return slide


def widget_text(slide, params: dict, values: dict, x, y, w, h):
    color = params.get("color") or _c("ink")
    text = fill(params.get("text") or "", values)
    paras = [(ln or " ", params.get("size") or 14, color, params.get("bold"), params.get("align") or "left")
             for ln in (text.splitlines() or [""])]
    _textbox(slide, x, y, w, h, paras, margin=0.12 if params.get("fill") else 0.0,
             fill_hex=params.get("fill") or None)


def widget_placeholder(slide, label: str, x, y, w, h):
    """What a widget without data shows: its name, in a dashed frame, never a gap."""
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.background(); sh.shadow.inherit = False
    sh.line.color.rgb = _rgb(_c("line")); sh.line.width = Pt(1)
    sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    tf = sh.text_frame
    r = tf.paragraphs[0].add_run(); r.text = label
    r.font.size = Pt(pptxtpl.font_size(11)); r.font.color.rgb = _rgb(_c("muted"))


# --------------------------------------------------------------------------
# A slide imported from a PowerPoint file
# --------------------------------------------------------------------------
_RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_A_T = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"


def import_slide(prs, blob: bytes, index: int, values: dict) -> tuple[object | None, list[str]]:
    """Copy slide ``index`` (1-based) of the PowerPoint ``blob`` into the deck: its
    shapes, its pictures and its background, every ``{field}`` of its text filled
    in. Charts, media and embedded objects are not carried (they belong to parts
    this copy does not bring): they are counted in the returned warnings."""
    from pptx import Presentation
    warnings: list[str] = []
    try:
        src = Presentation(io.BytesIO(blob))
    except Exception:
        return None, ["imported_unreadable"]
    slides = list(src.slides)
    if not slides:
        return None, ["imported_empty"]
    sl = slides[max(0, min(len(slides) - 1, int(index or 1) - 1))]
    slide = pptxtpl.add_slide(prs)
    tree = slide.shapes._spTree
    for el in sl.shapes._spTree.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag in ("nvGrpSpPr", "grpSpPr"):
            continue
        if tag == "graphicFrame" and el.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}tbl") is None:
            warnings.append("imported_object_skipped")
            continue
        new = copy.deepcopy(el)
        ok = True
        for blip in new.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
            rid = blip.get(_RID)
            if not rid:
                continue
            try:
                part = sl.part.related_part(rid)
                _img, new_rid = slide.part.get_or_add_image_part(io.BytesIO(part.blob))
                blip.set(_RID, new_rid)
            except Exception:
                ok = False
        if not ok:
            warnings.append("imported_object_skipped")
            continue
        for t in new.iter(_A_T):
            if t.text and "{" in t.text:
                t.text = fill(t.text, values)
        tree.append(new)
    bg = sl._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}cSld/"
                          "{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
    if bg is not None and bg.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip") is None:
        csld = slide._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}cSld")
        csld.insert(0, copy.deepcopy(bg))
    return slide, warnings


# --------------------------------------------------------------------------
# Footer and page numbers, laid on every slide after the deck is built
# --------------------------------------------------------------------------
COVER_NAME = "tfu-cover"


def footer(prs, text: str, page_numbers: bool, values: dict) -> None:
    """The document's footer and page numbers, on every slide but its covers (a
    grey line on a dark cover is unreadable, and a cover is not numbered)."""
    slides = list(prs.slides)
    n = len(slides)
    for i, slide in enumerate(slides, start=1):
        if slide._element.cSld.get("name") == COVER_NAME:
            continue
        vals = {**values, "page": f"{i}/{n}"}
        if text:
            _textbox(slide, 0.4, SH - 0.27, SW - 2.2, 0.22,
                     [(fill(text, vals), 8, _c("muted"), False, "left")], wrap=False)
        if page_numbers:
            _textbox(slide, SW - 1.6, SH - 0.27, 1.2, 0.22,
                     [(f"{i}/{n}", 8, _c("muted"), False, "right")], wrap=False)
