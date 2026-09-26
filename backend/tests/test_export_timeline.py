"""La frise annuelle telle qu'elle part: en HTML, en JPG et en PPTX.

Les trois exports racontaient l'annee en trois blocs separes, comme le faisait la
page: les initiatives, les engagements OTD, puis une roadmap en quatre colonnes.
Ils les posent maintenant sur un seul axe, le temps, qui est la seule chose que
ces objets ont en commun.

Ce qui se teste ici n'est pas la mise en page, qui bougera, mais ce qu'elle ne
doit jamais perdre: un jalon rattache a son engagement, un jalon rattache a rien
qui reste visible quand meme, un engagement pose au bon mois, un engagement sans
date qui est dit plutot que tu, et deux engagements qui ne se recouvrent pas. Le
JPG n'a pas de test propre: il est produit par le navigateur a partir de ce meme
HTML, donc le tester reviendrait a tester le navigateur.
"""
from __future__ import annotations

import datetime as dt
import io
import re
import zipfile

import pytest
from sqlalchemy import select

from app import report as report_mod
from app.models import (Initiative, KeyMessage, Objective, Otd, RoadmapItem, Squad,
                        SquadBudget, User)
from app.reportcommon import _MONTHS, pack_bands, timeline_groups

YEAR = 2026


def _utc(month: int, day: int, year: int = YEAR) -> dt.datetime:
    return dt.datetime(year, month, day, tzinfo=dt.timezone.utc)


@pytest.fixture()
def chained(db, seeded):
    """Une squad ou le chainage complet existe: une initiative, un objectif qui la
    sert, des jalons qui repondent a l'objectif, un jalon qui ne repond a rien, et
    des engagements OTD dates."""
    squad_id, tribe_id = seeded["squad_a"], seeded["t1"]
    squad = db.get(Squad, squad_id)
    squad.mood, squad.mood_at = "mixed", _utc(9, 1)
    squad.mood_comment = "Charge forte sur la fin de trimestre"

    init = Initiative(tribe_id=tribe_id, year=YEAR, title="Reduire le temps de build",
                      squad_id=squad_id, owner="Alice Martin", deadline=_utc(6, 30))
    db.add(init)
    db.flush()
    obj = Objective(squad_id=squad_id, year=YEAR, title="Diviser le demarrage par deux",
                    initiative_id=init.id, rag_status="amber", weight=1)
    db.add(obj)
    db.flush()
    leader = db.get(User, squad.leader_user_id)
    manage = Otd(tribe_id=tribe_id, year=YEAR, title="Plan VCF9 presente",
                 committed_date=_utc(2, 15), owner_user_id=leader.id, display_order=1)
    db.add(manage)
    db.add(Otd(tribe_id=tribe_id, year=YEAR, title="Adoption du modele dans deux pays",
               committed_date=_utc(2, 20), owner_user_id=leader.id, display_order=2))
    db.add(Otd(tribe_id=tribe_id, year=YEAR, title="Engagement sans date arretee",
               committed_date=None, owner_user_id=leader.id, display_order=3))
    # L'engagement que la squad prend elle-meme, a cote de celui du management.
    own = Otd(tribe_id=tribe_id, year=YEAR, title="Bascule 25G", scope="squad",
              squad_id=squad_id, committed_date=_utc(9, 30), owner_user_id=leader.id,
              display_order=4)
    db.add(own)
    db.flush()

    db.add(RoadmapItem(squad_id=squad_id, year=YEAR, quarter=1, title="Cache des dependances",
                       release_stage="EA", status="on_track", objective_id=obj.id,
                       initiative_id=init.id, otd_id=manage.id))
    # Celui-ci tient les deux engagements a la fois: le lien du management et
    # celui de la squad sont deux colonnes distinctes, justement pour cela.
    db.add(RoadmapItem(squad_id=squad_id, year=YEAR, quarter=3, title="Build incremental",
                       release_stage="GA", status="at_risk", objective_id=obj.id,
                       initiative_id=init.id, otd_id=manage.id, squad_otd_id=own.id))
    # Celui-la ne tient aucun engagement, et son theme est saisi a la main plutot
    # que porte par une initiative: les deux chemins doivent titrer une boite.
    db.add(RoadmapItem(squad_id=squad_id, year=YEAR, quarter=2, title="Nettoyage des images",
                       release_stage="EA", status="blocked", theme="Hygiene des images"))
    db.commit()
    return squad_id, db.scalar(select(User).where(User.email == seeded["admin"]))


def _data(db, squad_id, viewer, **kw):
    return report_mod.build_report_data(db, None, YEAR, 7, squad_id=squad_id,
                                        viewer=viewer, lang="fr", **kw)


def _deck_text(blob: bytes) -> str:
    runs: list[str] = []
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        for name in z.namelist():
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                runs += re.findall(r"<a:t>(.*?)</a:t>", z.read(name).decode("utf-8"), re.S)
    return "\n".join(runs)


# ----- ce que la frise assemble --------------------------------------------------

def _boxes(db, squad_id, viewer):
    det = _data(db, squad_id, viewer)["tribes"][0]["squads"][0]["detail"]
    return {g["title"]: g for g in timeline_groups(det)}


def test_a_box_carries_the_commitment_and_the_milestones_that_hold_it(db, chained):
    """Une boite porte un titre et une seule date, ce qui est la condition pour
    qu'un trait suffise a dire quand ses jalons sont attendus. Le titre est
    l'engagement qu'ils tiennent, la date est la sienne."""
    squad_id, viewer = chained
    box = _boxes(db, squad_id, viewer)["Plan VCF9 presente"]

    assert box["month"] == 1, "fevrier, la date de l'engagement"
    assert {it["title"] for it in box["items"]} == {"Cache des dependances",
                                                    "Build incremental"}


def test_a_milestone_holding_two_commitments_is_in_both_boxes(db, chained):
    """Le meme travail lu par deux promesses prises a deux dates. Le taire sous
    l'une reviendrait a dire que cette promesse n'existe pas."""
    squad_id, viewer = chained
    boxes = _boxes(db, squad_id, viewer)

    assert "Build incremental" in {it["title"] for it in boxes["Plan VCF9 presente"]["items"]}
    assert [it["title"] for it in boxes["Bascule 25G"]["items"]] == ["Build incremental"]
    assert boxes["Bascule 25G"]["month"] == 8, "septembre, l'engagement de la squad"


def test_a_milestone_appears_once_in_a_box(db, chained, seeded):
    """Un jalon qui tient deux fois le meme engagement ne s'y lit pas deux fois."""
    squad_id, viewer = chained
    squad = db.get(Squad, squad_id)
    second = Otd(tribe_id=seeded["t1"], year=YEAR, title="Second engagement de fevrier",
                 committed_date=_utc(2, 20), owner_user_id=squad.leader_user_id)
    db.add(second)
    db.flush()
    both = db.scalar(select(RoadmapItem).where(RoadmapItem.title == "Cache des dependances"))
    both.squad_otd_id = second.id
    db.commit()

    boxes = _boxes(db, squad_id, viewer)
    titles = [it["title"] for it in boxes["Second engagement de fevrier"]["items"]]
    assert titles == ["Cache des dependances"]


def test_the_two_scopes_are_told_apart_in_the_timeline_data(db, chained):
    """La distinction doit exister dans la donnee, pas seulement dans la couleur:
    c'est elle que les deux rendus peignent."""
    squad_id, viewer = chained
    det = _data(db, squad_id, viewer)["tribes"][0]["squads"][0]["detail"]
    scopes = {o["title"]: o["scope"] for o in det["otds"]}

    assert scopes["Plan VCF9 presente"] == "management"
    assert scopes["Bascule 25G"] == "squad"
    # Le management d'abord: la frise se lit de l'engagement subi vers celui qu'on choisit.
    assert [o["scope"] for o in det["otds"]][-1] == "squad"


def test_a_milestone_holding_nothing_keeps_a_box_of_its_own(db, chained):
    """Un jalon absent d'un export se lit comme un jalon qui n'existe pas. Sans
    engagement, sa boite est celle de son theme, et sa date le milieu de son
    trimestre: c'est la seule que la donnee porte."""
    squad_id, viewer = chained
    box = _boxes(db, squad_id, viewer)["Hygiene des images"]

    assert [it["title"] for it in box["items"]] == ["Nettoyage des images"]
    assert box["month"] == 4, "le milieu du Q2"
    assert box["dated"] is False


def test_the_commitments_carry_their_month_and_their_status(db, chained):
    """Le mois est ce qui pose un engagement sur l'axe; sans lui il n'a pas de place."""
    squad_id, viewer = chained
    det = _data(db, squad_id, viewer)["tribes"][0]["squads"][0]["detail"]
    by_title = {o["title"]: o for o in det["otds"]}

    assert by_title["Plan VCF9 presente"]["month"] == 1          # fevrier
    assert by_title["Engagement sans date arretee"]["month"] is None
    assert by_title["Plan VCF9 presente"]["status"] == "late"    # date passee, rien de livre


def test_a_commitment_of_another_year_does_not_land_on_this_axis(db, chained, seeded):
    """Une date de 2025 n'a pas de mois sur la frise 2026: elle serait posee a faux."""
    squad_id, viewer = chained
    squad = db.get(Squad, squad_id)
    db.add(Otd(tribe_id=seeded["t1"], year=YEAR, title="Reste de l'an dernier",
               committed_date=_utc(11, 30, YEAR - 1), owner_user_id=squad.leader_user_id))
    db.commit()

    det = _data(db, squad_id, viewer)["tribes"][0]["squads"][0]["detail"]
    assert next(o for o in det["otds"] if o["title"] == "Reste de l'an dernier")["month"] is None


def _span(chars_per_month=8):
    """La place que prend un libelle, en colonnes, comme le font les deux rendus."""
    return lambda it: max(1, -(-len(it["title"]) // chars_per_month))


def test_two_markers_close_together_are_stacked_not_superimposed(db, chained):
    """Deux engagements de fevrier se genent: l'un descend d'une bande."""
    squad_id, viewer = chained
    det = _data(db, squad_id, viewer)["tribes"][0]["squads"][0]["detail"]
    placed, hidden = pack_bands(det["otds"], _span())

    assert hidden == 0
    assert {row for _, _, _, row in placed} == {0, 1}
    spans = [(start, start + width, row) for _, start, width, row in placed]
    assert all(not (a0 < b1 and b0 < a1) for i, (a0, a1, ar) in enumerate(spans)
               for (b0, b1, br) in spans[i + 1:] if ar == br)


def test_a_long_title_takes_more_room_than_a_short_one(db):
    """Un titre tronque a trois lettres ne dit rien: la place reservee suit le
    texte. Une reserve fixe laissait en plus un blanc derriere les libelles
    courts, et poussait leurs voisins une bande plus bas pour rien."""
    items = [{"title": "Court", "month": 0},
             {"title": "Un titre nettement plus long", "month": 6}]
    placed, _ = pack_bands(items, _span())
    widths = {it["title"]: w for it, _, w, _ in placed}
    assert widths["Un titre nettement plus long"] > widths["Court"]


def test_a_december_label_is_written_to_the_left_of_its_marker(db):
    """Un repere de decembre n'a plus rien devant lui: son libelle s'ecrit a
    gauche, et le bloc recule juste assez pour le loger. Cale sur la fin de
    l'annee, il reservait des colonnes a droite du repere, la ou rien ne s'ecrit."""
    items = [{"title": "Bascule complete du parc en production", "month": 11}]
    (it, start, width, _row), = pack_bands(items, _span())[0]

    assert start < it["month"], "le bloc de titre commence avant decembre"
    assert start + width == it["month"] + 1, "et s'arrete juste apres le repere"


def test_a_band_that_is_full_says_so_instead_of_dropping_the_rest(db):
    """Sur une slide la hauteur ne s'etire pas; ce qui ne tient pas est compte."""
    items = [{"title": f"Engagement {i}", "month": 3} for i in range(6)]
    placed, hidden = pack_bands(items, _span(16), max_rows=3)
    assert len(placed) == 3 and hidden == 3


# ----- ce que les documents montrent ---------------------------------------------

def test_the_html_export_shows_the_whole_block(db, chained):
    squad_id, viewer = chained
    html = report_mod.render_html(_data(db, squad_id, viewer))

    assert 'class="xtl-grid"' in html
    assert len(re.findall(r"<b>Q[1-4]</b>", html)) == 4, "les quatre trimestres"
    assert len(re.findall(r'class="xtl-m"', html)) == 12, "les douze mois"
    assert "Plan VCF9 presente" in html and "Cache des dependances" in html
    # Le jalon est la; la boite qui le porte n'annonce aucune categorie.
    assert "Nettoyage des images" in html
    assert "hors engagement" not in html
    # L'etoile est dessinee, pas prise dans une police, et la boite est en pointille.
    assert "xtl-star" in html and "border:1px dashed" in html
    # Et la boite porte le titre de l'engagement, ou le theme a defaut.
    assert "Plan VCF9 presente" in html and "Hygiene des images" in html
    # Les deux portees se distinguent, et leur legende voyage avec la frise.
    assert "xtl-scope-squad" in html and "Engagement de la squad" in html
    # Un engagement sans date est cite sous la bande, pas efface.
    assert "Engagements sans date" in html and "Engagement sans date arretee" in html
    # Le moral, avec la date qui le date (le libelle part echappe, pas le niveau).
    assert 'class="xtl-mood"' in html and "Moyen" in html and "01/09/2026" in html
    # La feuille de la frise voyage avec elle: l'export JPG n'en charge pas d'autre.
    assert ".xtl-box" in html


def test_the_deck_puts_one_squad_on_one_slide(db, chained):
    squad_id, viewer = chained
    deck = report_mod.render_pptx(_data(db, squad_id, viewer))
    text = _deck_text(deck)

    assert "Frise 2026" in text
    assert all(f"Q{q}" in text for q in (1, 2, 3, 4))
    assert "Build incremental" in text
    assert "Nettoyage des images" in text
    assert "Plan VCF9 presente" in text and "Bascule 25G" in text
    assert "Engagement management" in text, "la legende des deux portees"
    # Le moral: « Moral » au-dessus, « équipe » en dessous, et le nuage a droite. Le
    # niveau n'est plus ecrit a cote: il se lit sur le visage et sur la couleur,
    # et l'ecrire revenait a le dire deux fois dans deux centimetres.
    assert "Moral" in text and "équipe" in text
    pptx = pytest.importorskip("pptx")
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches
    prs = pptx.Presentation(io.BytesIO(deck))
    # Le nuage est dessine point par point, et non pris dans les formes de
    # PowerPoint, dont le nuage a d'autres bosses que celui de l'application. On
    # le reconnait a sa taille et a sa place: en haut a droite, dans la carte du
    # moral. La bouche est un second contour, plus petit, dans le meme coin.
    clouds = [sh for sh in prs.slides[0].shapes
              if sh.shape_type == MSO_SHAPE_TYPE.FREEFORM
              and sh.top < Inches(1.2) and sh.width > Inches(0.3)]
    assert len(clouds) == 1, "un nuage, celui du moral declare"


def test_the_deck_says_the_release_stage_of_each_milestone(db, chained):
    """EA ou GA repond a « est-ce ouvert a tout le monde », qui est la premiere
    question posee sur un jalon en comite. Les deux lettres sont dans la boite du
    jalon, et leur legende est au bas de la slide: une abreviation sans legende ne
    dit rien a qui decouvre le document."""
    pptx = pytest.importorskip("pptx")
    squad_id, viewer = chained
    blob = report_mod.render_pptx(_data(db, squad_id, viewer))
    text = _deck_text(blob)
    assert "EA" in text and "GA" in text
    assert "Accès anticipé" in text and "Disponibilité générale" in text

    # La phase suit son jalon dans la meme puce, a l'interieur de la carte: elle
    # ne peut donc plus se poser sur le titre, ce qui etait le risque du modele
    # precedent, ou les deux vivaient dans deux formes superposees.
    prs = pptx.Presentation(io.BytesIO(blob))
    cards = [sh for sh in prs.slides[0].shapes
             if sh.has_text_frame and "Build incremental" in sh.text_frame.text]
    assert cards, "le jalon est dans une carte"
    for card in cards:
        assert "Build incremental (GA)" in card.text_frame.text
        assert card.text_frame.word_wrap is True, "le texte passe a la ligne, il ne se coupe pas"


def test_a_box_is_titled_and_that_title_is_on_the_slide(db, chained):
    """Une boite doit dire de quoi elle parle: elle est rangee dans l'ordre des
    dates et non sous son etoile, donc son trait seul dirait « cette date-la »
    sans dire « cette promesse-la ». Le titre est l'engagement que ses jalons
    tiennent, ou le theme quand ils n'en tiennent aucun (l'initiative qu'ils
    servent, ou le theme que leur squad leader a saisi)."""
    pptx = pytest.importorskip("pptx")
    squad_id, viewer = chained
    assert set(_boxes(db, squad_id, viewer)) == {"Plan VCF9 presente", "Bascule 25G",
                                                 "Hygiene des images"}

    prs = pptx.Presentation(io.BytesIO(report_mod.render_pptx(_data(db, squad_id, viewer))))
    names = {sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame}
    assert {"Plan VCF9 presente", "Bascule 25G", "Hygiene des images"} <= names


def test_no_title_is_cut_in_the_deck(db, chained):
    """La demande tient en une phrase: on voit tout le texte, la largeur est fixe,
    et au pire ca passe a la ligne. Un titre coupe ne dit plus rien, et personne
    ne peut deviner ce qui manque."""
    squad_id, viewer = chained
    text = _deck_text(report_mod.render_pptx(_data(db, squad_id, viewer)))

    for title in ("Cache des dependances", "Build incremental", "Nettoyage des images",
                  "Plan VCF9 presente", "Adoption du modele dans deux pays", "Bascule 25G"):
        assert title in text, f"« {title} » n'est pas sorti en entier"
    assert "\u2026" not in text, "aucun titre n'est tronque par des points de suspension"


def test_no_shape_falls_outside_the_slide(db, chained):
    """Une forme hors cadre ne se voit pas a la lecture du code, seulement a la
    projection."""
    pptx = pytest.importorskip("pptx")
    squad_id, viewer = chained
    blob = report_mod.render_pptx(_data(db, squad_id, viewer))
    prs = pptx.Presentation(io.BytesIO(blob))
    w, h = prs.slide_width, prs.slide_height

    outside = [(s_i, sh.name) for s_i, s in enumerate(prs.slides) for sh in s.shapes
               if sh.left < 0 or sh.top < 0 or sh.left + sh.width > w or sh.top + sh.height > h]
    assert not outside, outside[:5]


def test_a_quarter_card_covers_exactly_its_three_months(db, chained):
    """Les deux bandes se lisent ensemble, donc elles sont coupees au meme endroit.

    La carte d'un trimestre prenait toute sa gouttiere a droite: elle etait plus
    courte que ses trois mois et s'arretait avant la fin du troisieme. Et la bande
    des mois etait d'un seul tenant sous des cartes separees, donc rien n'y disait
    ou un trimestre finissait: mars semblait deborder de Q1. L'axe est desormais
    decoupe par trimestre des deux cotes, avec la meme gouttiere.
    """
    pptx = pytest.importorskip("pptx")
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    squad_id, viewer = chained
    prs = pptx.Presentation(io.BytesIO(report_mod.render_pptx(_data(db, squad_id, viewer))))
    slide = list(prs.slides)[-1]

    cards = sorted((sh for sh in slide.shapes
                    if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                    and sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
                    and sh.has_text_frame and sh.text_frame.text.startswith("Q")),
                   key=lambda sh: sh.left)
    assert len(cards) == 4, [sh.text_frame.text for sh in cards]
    months = sorted((sh for sh in slide.shapes
                     if sh.has_text_frame
                     and sh.text_frame.text in _MONTHS["fr"]), key=lambda sh: sh.left)
    assert len(months) == 12, [sh.text_frame.text for sh in months]

    for i, card in enumerate(cards):
        first, last = months[i * 3], months[i * 3 + 2]
        assert card.left == first.left, f"Q{i + 1} ne commence pas au bord de son premier mois"
        assert card.left + card.width == last.left + last.width, (
            f"Q{i + 1} ne finit pas au bord de son troisieme mois")
    # Et la coupure est la meme dans les deux bandes: la gouttiere qui separe deux
    # cartes de trimestre separe aussi leurs deux blocs de mois, au meme endroit.
    for i, (a, b) in enumerate(zip(cards, cards[1:])):
        last, first = months[i * 3 + 2], months[i * 3 + 3]
        assert (a.left + a.width, b.left) == (last.left + last.width, first.left), (
            f"la coupure Q{i + 1}/Q{i + 2} ne tombe pas au meme endroit dans les deux bandes")
        assert b.left > a.left + a.width, "les deux blocs doivent etre separes"


def test_a_quarter_title_never_sits_on_its_own_progress_bar(db, chained):
    """Le titre d'un trimestre et sa barre partagent une carte d'un demi-pouce.

    La barre etait posee a une distance fixe du haut de la carte, calee sur un
    corps de 13 points. Le corps ayant grandi, la ligne de titre descend plus bas
    que cette distance: « Q4  50 % » s'ecrivait par-dessus sa propre barre. La
    hauteur de la carte et la position de la barre se deduisent donc du corps
    reel, et ce test mesure la ligne plutot que de faire confiance au chiffre.
    """
    pptx = pytest.importorskip("pptx")
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    squad_id, viewer = chained
    prs = pptx.Presentation(io.BytesIO(report_mod.render_pptx(_data(db, squad_id, viewer))))
    slide = list(prs.slides)[-1]

    rounded = [sh for sh in slide.shapes
               if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
               and sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE]
    cards = [sh for sh in rounded if sh.has_text_frame and sh.text_frame.text.startswith("Q")]
    assert len(cards) == 4, [sh.text_frame.text for sh in cards]

    for card in cards:
        # La barre est le rectangle arrondi sans texte pose dans cette carte.
        bars = [b for b in rounded
                if not (b.has_text_frame and b.text_frame.text.strip())
                and card.left <= b.left and b.left + b.width <= card.left + card.width
                and card.top <= b.top and b.top + b.height <= card.top + card.height]
        if "rien de prévu" in card.text_frame.text or "nothing planned" in card.text_frame.text:
            # A quarter with nothing planned has no bar on purpose (it is not 0 %).
            assert not bars
            continue
        assert bars, f"« {card.text_frame.text} » a perdu sa barre d'avancement"
        run = card.text_frame.paragraphs[0].runs[0]
        line = int(run.font.size.pt * 1.2 / 72 * 914400)     # une ligne vaut 1,2 corps
        text_bottom = card.top + (card.text_frame.margin_top or 0) + line
        assert min(b.top for b in bars) >= text_bottom, (
            f"« {card.text_frame.text} » s'ecrit sur sa barre")
        assert max(b.top + b.height for b in bars) <= card.top + card.height, (
            f"la barre de « {card.text_frame.text} » sort de sa carte")


def test_the_bottom_cards_are_tall_enough_for_their_own_text(db, chained):
    """Messages cles et budget tiennent dans leur encadre.

    La carte du budget porte un titre et trois lignes, et elle etait plus courte
    que cette hauteur: « Prevision » sortait par le bas et se lisait a moitie. Rien
    dans le code ne le disait, et le corps du texte ayant grossi depuis, la ligne
    passait sur la legende de la slide.
    """
    pptx = pytest.importorskip("pptx")
    from pptx.util import Pt
    squad_id, viewer = chained
    # Les deux cartes remplies: sans budget ni message, elles ne portent qu'une
    # ligne et ne disent rien de leur hauteur.
    db.get(Squad, squad_id).budget_enabled = True
    db.add(SquadBudget(squad_id=squad_id, year=YEAR, total=100000, spent=80000,
                       forecast=125000))
    for kind, text in (("risk", "Le fournisseur a glisse d'une semaine"),
                       ("success", "Les equipes pilotes sont en production")):
        db.add(KeyMessage(squad_id=squad_id, year=YEAR, kind=kind, text=text))
    db.commit()
    blob = report_mod.render_pptx(_data(db, squad_id, viewer))
    assert "125\u202f000" in _deck_text(blob), (
        "la troisieme ligne de budget doit etre rendue pour que la mesure ait un sens")
    prs = pptx.Presentation(io.BytesIO(blob))
    slide = list(prs.slides)[-1]

    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    short = []
    for sh in slide.shapes:
        # Les deux encadres du bas, et eux seuls: c'est le cadre qui rogne son
        # texte, une zone de texte libre deborde sans se voir.
        if (sh.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE
                or sh.auto_shape_type != MSO_SHAPE.ROUNDED_RECTANGLE
                or not sh.has_text_frame or not sh.text_frame.text.strip()
                or sh.top < Pt(6.2 * 72)):
            continue
        tf = sh.text_frame
        # Une ligne vaut 1,2 fois son corps, plus l'espace qui la suit.
        need = sum(max((r.font.size.pt for r in para.runs if r.font.size), default=0) * 1.2
                   + (para.space_after.pt if para.space_after else 0)
                   for para in tf.paragraphs if para.runs)
        need += (tf.margin_top.pt or 0) + (tf.margin_bottom.pt or 0)
        if need > sh.height / 12700:
            short.append((tf.text[:40], round(need, 1), round(sh.height / 12700, 1)))
    assert not short, short


def test_the_multi_squad_deck_keeps_its_summary_then_one_slide_per_squad(db, chained, seeded):
    """Un export de plusieurs squads garde la synthese; un export d'une seule squad
    est le meme deck sans elle."""
    pptx = pytest.importorskip("pptx")
    _, viewer = chained
    wide = report_mod.build_report_data(db, seeded["t1"], YEAR, 7, viewer=viewer, lang="fr")
    n = len(wide["tribes"][0]["squads"])

    prs = pptx.Presentation(io.BytesIO(report_mod.render_pptx(wide)))
    # La synthese, une slide par squad, puis les points d'attention du rapport hebdo.
    assert len(prs.slides) == n + 2

    one = _data(db, seeded["squad_a"], viewer)
    assert len(pptx.Presentation(io.BytesIO(report_mod.render_pptx(one))).slides) == 1


def test_a_squad_with_nothing_says_so_rather_than_showing_an_empty_grid(db, seeded):
    """Une squad vierge est un cas normal en janvier, pas une erreur."""
    viewer = db.scalar(select(User).where(User.email == seeded["admin"]))
    html = report_mod.render_html(_data(db, seeded["squad_b"], viewer))
    assert "Aucun engagement, aucun jalon" in html
    assert "Aucun OTD" in html
