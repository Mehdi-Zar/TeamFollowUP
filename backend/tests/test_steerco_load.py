"""Steerco PowerPoint under load: whatever the number of KPIs, SLA services and
events, every shape stays on the slide and nothing is silently dropped.

A platform slide is one page with fixed panels. Past what a panel holds, the
KPIs and the SLA turn into lists, and what the lists cannot hold, like the
events past their panel, goes in full to a detail slide right after it.
"""
import io
import os

import pytest

from app.models import Platform, SteercoEntry, Squad
from app.platforms import default_template
from tests.conftest import login
from tests.test_steerco import _enable

PERIOD = "2026-06"
LONG = "Migration des environnements historiques vers la nouvelle plateforme avec bascule progressive"


def _platform(db, seeded, n_kpis, n_sla, n_events):
    sid = seeded["squad_a"]
    tpl = default_template(sid)
    tpl["kpis"] += [{"label": f"Indicateur supplementaire numero {k + 1}", "owner_squad_id": sid, "sub": []}
                    for k in range(max(0, n_kpis - len(tpl["kpis"])))]
    tpl["sla"] += [{"label": f"Service critique {k + 1}", "owner_squad_id": sid}
                   for k in range(max(0, n_sla - len(tpl["sla"])))]
    p = Platform(tribe_id=seeded["t1"], name="Plateforme de charge", display_order=1,
                 steerco_enabled=True, template=tpl)
    p.contributors = [db.get(Squad, sid)]
    db.add(p)
    db.flush()
    ev = [{"date": f"{d:02d}/06", "text": f"Evenement {d} {LONG}", "tag": "Incident",
           "sev": ["red", "amber", "green", "ice"][d % 4], "squad_id": sid} for d in range(1, n_events + 1)]
    for mm in range(1, 7):
        db.add(SteercoEntry(platform_id=p.id, period=f"2026-{mm:02d}", data={
            "kpis": [{"label": k["label"], "value": str(1000 + 37 * mm + i)} for i, k in enumerate(tpl["kpis"])],
            "sla": {"services": [x["label"] for x in tpl["sla"]],
                    "cells": [{"v": str(80 + (i + mm) % 20)} for i, _ in enumerate(tpl["sla"])]},
            "incidents": str(mm),
            "last_events": ev if mm == 6 else [], "next_events": ev if mm == 6 else []}))
    db.commit()
    return ev


def _deck(client, lang="fr"):
    from pptx import Presentation
    r = client.get(f"/api/steerco/document.pptx?period={PERIOD}&lang={lang}")
    assert r.status_code == 200, r.text
    if os.environ.get("STEERCO_DUMP"):
        open(os.environ["STEERCO_DUMP"], "wb").write(r.content)
    return Presentation(io.BytesIO(r.content))


def _text(prs):
    out = []
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            elif sh.has_table:
                out += [c.text for row in sh.table.rows for c in row.cells]
    return "\n".join(out)


def _off_slide(prs):
    W, H = prs.slide_width, prs.slide_height
    return [(i, sh.name) for i, s in enumerate(prs.slides) for sh in s.shapes
            if sh.left < 0 or sh.top < 0 or sh.left + sh.width > W + 9525 or sh.top + sh.height > H + 9525]


@pytest.mark.parametrize("n_kpis,n_sla,n_events", [(5, 3, 2), (8, 6, 6), (13, 12, 18), (40, 30, 25)])
def test_every_item_fits_or_goes_to_the_detail_slide(client, db, seeded, n_kpis, n_sla, n_events):
    pytest.importorskip("pptx")
    _enable(client)
    ev = _platform(db, seeded, n_kpis, n_sla, n_events)
    prs = _deck(client)
    assert not _off_slide(prs)
    text = " ".join(_text(prs).split())
    heavy = n_events > 3
    assert ("détail" in text) == heavy
    # The light case keeps its one slide; under load every event reads in full.
    assert len(prs.slides) == 1 if not heavy else len(prs.slides) > 1
    if heavy:
        for e in ev:
            assert e["text"] in text, e["text"]
        assert "en slide suivante" in text
    if n_kpis > 20:
        # On the slide (cut in the middle, the end kept) or on the detail slide.
        base = default_template(None)
        assert f'numero {n_kpis - len(base["kpis"])}' in text
        assert f'critique {n_sla - len(base["sla"])}' in text
    # Labels that share a long start stay distinguishable on the slide.
    if n_kpis > 10:
        first = prs.slides[0]
        labels = [c.text for sh in first.shapes if sh.has_table for row in sh.table.rows for c in row.cells]
        assert not any(t.endswith("Indicateur…") or t == "Indicateur…" for t in labels)
