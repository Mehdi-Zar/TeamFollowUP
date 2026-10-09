"""From the reporting to the inbox, end to end.

A squad leader fills the weekly reporting (milestones with dependencies, a
commitment, KPIs, key messages, mood, the quarter's comment) and submits it. Then
every document the users take away is produced, and every mail the application
sends leaves through a real SMTP connection to a local server that keeps what it
receives. Each document and each mail is read back the way its reader gets it:
the PPTX opened, the HTML reduced to its visible text, the mail parsed as a mail
client would, its attachments opened.

What is checked everywhere:
- the figures typed in the reporting are the ones in the documents and the mails;
- no placeholder or leaked value reaches the reader (None, undefined, NaN, "{x}");
- no em dash or middle dot (the project's typography rule);
- a mail has a subject, a sender, a date, an id, an HTML and a text part, and its
  attachments open.

Set E2E_OUT to a folder to also write every document and mail there, for a look.
"""
import datetime as dt
import email
import io
import os
import re
import socket
from email import policy
from html import unescape

import pytest

import app.changenotify as changenotify
import app.database as dbmod
from app import report as report_mod
from app.changeconfig import ALL_EVENTS, set_change_notify
from app.reportconfig import set_report
from app.smtpconfig import set_smtp
from tests.conftest import TestingSessionLocal, login

YEAR = dt.datetime.now(dt.timezone.utc).year
OUT = os.environ.get("E2E_OUT")

# What the squad leader types; every one of these must reach the documents.
MILESTONE_BLOCKED = "Migration Vault vers HSM managé"
MILESTONE_ON_TRACK = "Portail self-service des landing zones"
DEP_TEXT_VENDOR = "Livraison firewall par le fournisseur réseau"
OTD_TITLE = "Kubernetes managé ouvert"
KEY_MESSAGE = "Premier client migré sur le socle sans incident"
QUARTER_COMMENT = "Trimestre tenu malgré le retard du HSM"
KPI_NAME = "Clusters en production"

LEAKS = re.compile(r"\bNone\b|\bundefined\b|\bNaN\b|\bnull\b|\{[a-z_]+\}|\[object Object\]")
BANNED = (chr(0x2014), chr(0x00B7))   # em dash, middle dot


# --------------------------------------------------------------------------- SMTP sink
class _Sink:
    """A real SMTP server on a free local port, keeping every message it gets."""

    def __init__(self):
        from aiosmtpd.controller import Controller
        self.messages: list[email.message.EmailMessage] = []
        self.envelopes: list[tuple[str, list[str]]] = []
        sink = self

        class Handler:
            async def handle_DATA(self, server, session, envelope):
                sink.envelopes.append((envelope.mail_from, list(envelope.rcpt_tos)))
                sink.messages.append(email.message_from_bytes(envelope.content, policy=policy.default))
                return "250 OK"

        with socket.socket() as s:
            s.bind(("127.0.0.1", 0))
            self.port = s.getsockname()[1]
        self.controller = Controller(Handler(), hostname="127.0.0.1", port=self.port)
        self.controller.start()

    def stop(self):
        self.controller.stop()

    def take(self) -> list:
        out, self.messages = self.messages, []
        return out


@pytest.fixture
def inbox():
    sink = _Sink()
    yield sink
    sink.stop()


# --------------------------------------------------------------------------- readers
def _visible_text(html: str) -> str:
    html = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", html)
    html = re.sub(r"(?s)<!--.*?-->", " ", html)
    return re.sub(r"\s+", " ", unescape(re.sub(r"<[^>]+>", " ", html))).strip()


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    parts = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                parts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                parts += [c.text for row in sh.table.rows for c in row.cells]
            if getattr(sh, "has_chart", False) and sh.has_chart:
                parts += [s.name for s in sh.chart.plots[0].series] if sh.chart.plots else []
    return "\n".join(parts)


def _clean(text: str, where: str) -> None:
    leak = LEAKS.search(text)
    assert not leak, f"{where}: leaked value {leak.group(0)!r} near {text[max(0, leak.start() - 60):leak.end() + 60]!r}"
    for ch in BANNED:
        assert ch not in text, f"{where}: banned character U+{ord(ch):04X}"


def _save(name: str, data) -> None:
    if not OUT:
        return
    os.makedirs(OUT, exist_ok=True)
    mode = "wb" if isinstance(data, (bytes, bytearray)) else "w"
    with open(os.path.join(OUT, name), mode, **({} if mode == "wb" else {"encoding": "utf-8"})) as f:
        f.write(data)


def _read_mail(msg, where: str) -> dict:
    """A mail as its reader gets it: headers, text, HTML, attachments. Checked."""
    for h in ("From", "To", "Subject", "Date", "Message-ID"):
        assert msg[h], f"{where}: header {h} missing"
    subject = str(msg["Subject"])
    _clean(subject, f"{where} subject")
    html_part = msg.get_body(preferencelist=("html",))
    text_part = msg.get_body(preferencelist=("plain",))
    assert html_part is not None, f"{where}: no HTML part"
    assert text_part is not None, f"{where}: no text part"
    html = html_part.get_content()
    text = text_part.get_content()
    assert len(text.strip()) > 40, f"{where}: text part empty"
    _clean(_visible_text(html), f"{where} html")
    _clean(text, f"{where} text")
    atts = {}
    for part in msg.iter_attachments():
        name = part.get_filename()
        data = part.get_payload(decode=True)
        assert name and data, f"{where}: unnamed or empty attachment"
        atts[name] = data
        if name.endswith(".pptx"):
            _clean(_pptx_text(data), f"{where} attachment {name}")
        if name.endswith(".html"):
            _clean(_visible_text(data.decode("utf-8")), f"{where} attachment {name}")
    slug = re.sub(r"[^A-Za-z0-9]+", "_", where)[:60]
    _save(f"mail_{slug}.eml", msg.as_bytes())
    _save(f"mail_{slug}.html", html)
    for name, data in atts.items():
        _save(f"mail_{slug}__{name}", data)
    return {"subject": subject, "html": html, "visible": _visible_text(html), "text": text,
            "atts": atts, "to": str(msg["To"]), "cc": str(msg["Cc"] or "")}


# --------------------------------------------------------------------------- scenario
def _enable_everything(client):
    login(client, "admin@test")
    r = client.put("/api/admin/modules-config", json={
        "steerco": {"enabled": True},
        "squad_content": {"enabled": True, "roadmap": True, "kpis": True, "quarter_progress": True},
        "review": {"enabled": True, "weekly_report": True},
        "notifications": {"enabled": True, "inapp": True, "email": True},
    })
    assert r.status_code == 200, r.text


def _fill_reporting(client, seeded) -> dict:
    """The squad leader's week, through the same API calls as the screens."""
    sa, sb, t2 = seeded["squad_a"], seeded["squad_b"], seeded["t2"]
    login(client, seeded["sl_a"])
    ids = {}
    r = client.post("/api/roadmap-items", json={
        "squad_id": sa, "year": YEAR, "quarter": 2, "title": MILESTONE_BLOCKED, "theme": "Sécurité",
        "release_stage": "GA", "status": "blocked", "owner": "Claire Dubois",
        "dependency_kind": "squad", "dependency_squad_id": sb})
    assert r.status_code == 201, r.text
    ids["blocked"] = r.json()["id"]
    r = client.post("/api/roadmap-items", json={
        "squad_id": sa, "year": YEAR, "quarter": 3, "title": MILESTONE_ON_TRACK, "theme": "Plateforme",
        "release_stage": "EA", "status": "on_track", "owner": "Karim Benali",
        "dependency_kind": "tribe", "dependency_tribe_id": t2})
    assert r.status_code == 201, r.text
    ids["on_track"] = r.json()["id"]
    r = client.post("/api/roadmap-items", json={
        "squad_id": sa, "year": YEAR, "quarter": 4, "title": "Durcissement CIS des images", "theme": "Sécurité",
        "release_stage": "EA", "status": "at_risk", "dependency_kind": "text", "dependencies": DEP_TEXT_VENDOR})
    assert r.status_code == 201, r.text
    # The squad's own commitment, held by the blocked milestone.
    r = client.post("/api/otds", json={
        "tribe_id": seeded["t1"], "year": YEAR, "scope": "squad", "squad_id": sa, "title": OTD_TITLE,
        "committed_date": f"{YEAR}-06-30T00:00:00Z"})
    assert r.status_code == 201, r.text
    otd = r.json()["id"]
    assert client.put(f"/api/otds/{otd}/jalons", json={"jalon_ids": [ids["blocked"]]}).status_code == 200
    r = client.post("/api/kpis", json={"squad_id": sa, "name": KPI_NAME, "unit": "", "target_value": 12,
                                        "current_value": 9, "trend_status": "under_pressure"})
    assert r.status_code == 201, r.text
    r = client.post(f"/api/squads/{sa}/key-messages?year={YEAR}", json={"kind": "success", "text": KEY_MESSAGE})
    assert r.status_code == 201, r.text
    assert client.put(f"/api/squads/{sa}/mood", json={"mood": "mixed", "comment": "Équipe fatiguée par l'astreinte"}).status_code == 200
    r = client.put(f"/api/squads/{sa}/quarter-progress?year={YEAR}", json={"year": YEAR, "quarter": 2, "comment": QUARTER_COMMENT})
    assert r.status_code in (200, 201), r.text
    r = client.post(f"/api/squads/{sa}/snapshots", json={"year": YEAR})
    assert r.status_code in (200, 201), r.text
    return ids


def _steerco(client, db, seeded):
    from app.models import Platform, Squad
    from app.platforms import default_template
    p = Platform(tribe_id=seeded["t1"], name="Cloud Platform", display_order=1, steerco_enabled=True,
                 template=default_template(seeded["squad_a"]))
    p.contributors = [db.get(Squad, seeded["squad_a"])]
    db.add(p)
    db.commit()
    login(client, seeded["sl_a"])
    period = f"{YEAR}-06"
    r = client.put(f"/api/steerco/platform/{p.id}?period={period}", json={
        "kpis": [{"label": "Cloud Users", "value": "760"}, {"label": "Landing Zone", "value": "240"},
                 {"label": "K8aaS", "value": "4"}, {"label": "DBaaS", "value": "29"},
                 {"label": "Software Factory", "value": "5060"}],
        "sla": {"services": ["Incidents"], "cells": [{"v": "99,4%"}]},
        "incidents": "3",
        "last_events": [{"date": "12/06", "tag": "MEP", "text": "Mise en production du portail"}],
        "next_events": [{"date": "03/07", "tag": "COPIL", "text": "Revue trimestrielle"}]})
    assert r.status_code == 200, r.text
    return p.id, period


# --------------------------------------------------------------------------- the test
def test_from_the_reporting_to_the_documents(client, db, seeded):
    """Every export a user takes away carries what was typed, and nothing else."""
    pytest.importorskip("pptx")
    _enable_everything(client)
    _fill_reporting(client, seeded)
    pid, period = _steerco(client, db, seeded)
    sa = seeded["squad_a"]

    login(client, seeded["tribe"])
    docs = {
        "weekly": f"/api/reports/weekly.{{ext}}?year={YEAR}",
        "weekly_squad": f"/api/reports/weekly.{{ext}}?year={YEAR}&squad_id={sa}",
        "dashboard": f"/api/reports/dashboard.{{ext}}?year={YEAR}",
        "roadmap": f"/api/reports/roadmap.{{ext}}?year={YEAR}",
        "dependencies": f"/api/reports/dependencies.{{ext}}?year={YEAR}",
        "squad_roadmap": f"/api/squads/{sa}/roadmap.{{ext}}?year={YEAR}",
        "steerco": f"/api/steerco/document.{{ext}}?period={period}",
        "steerco_platform": f"/api/steerco/document.{{ext}}?period={period}&platform_id={pid}",
        "initiatives": f"/api/initiatives/report.{{ext}}?year={YEAR}",
        "org": "/api/org/export.{ext}",
    }
    expect = {
        "weekly": [MILESTONE_BLOCKED, "Squad A"],
        "weekly_squad": [MILESTONE_BLOCKED, MILESTONE_ON_TRACK, KEY_MESSAGE, KPI_NAME],
        "dashboard": ["Squad A"],
        "roadmap": [MILESTONE_BLOCKED, MILESTONE_ON_TRACK],
        "dependencies": [MILESTONE_BLOCKED, "Squad B"],
        "squad_roadmap": [MILESTONE_BLOCKED, MILESTONE_ON_TRACK],
        "steerco": ["Cloud Platform", "5060", "760"],
        "steerco_platform": ["Cloud Platform", "Mise en production du portail"],
        "initiatives": [],
        "org": [],
    }
    problems = []
    for name, url in docs.items():
        for ext in ("html", "pptx"):
            for lang in ("fr", "en"):
                full = url.format(ext=ext) + ("&" if "?" in url else "?") + f"lang={lang}"
                r = client.get(full)
                if r.status_code != 200:
                    problems.append(f"{name}.{ext} ({lang}): HTTP {r.status_code} {r.text[:120]}")
                    continue
                where = f"{name}.{ext} ({lang})"
                try:
                    if ext == "pptx":
                        assert r.content[:2] == b"PK", "not a zip"
                        text = _pptx_text(r.content)
                        _save(f"doc_{name}_{lang}.pptx", r.content)
                    else:
                        text = _visible_text(r.text)
                        _save(f"doc_{name}_{lang}.html", r.text)
                    _clean(text, where)
                    missing = [w for w in expect[name] if w not in text]
                    assert not missing, f"missing {missing}"
                except AssertionError as e:
                    problems.append(f"{where}: {e}")
    assert not problems, "\n".join(problems)


def test_from_the_reporting_to_the_inbox(client, db, seeded, inbox, monkeypatch):
    """Every mail leaves through SMTP and arrives readable, with its documents."""
    pytest.importorskip("pptx")
    _enable_everything(client)
    _fill_reporting(client, seeded)
    sa = seeded["squad_a"]
    login(client, "admin@test")
    set_smtp(db, {"enabled": True, "host": "127.0.0.1", "port": inbox.port, "use_tls": False,
                  "from_addr": "teamfollowup@example.org", "from_name": "TeamFollowUP"})
    from app.authconfig import set_auth_config
    set_auth_config(db, {"public_base_url": "https://teamfollowup.example.org"})
    db.commit()
    mails = {}

    # 1. "Send the report" from the export menu, to oneself.
    login(client, seeded["tribe"])
    r = client.post("/api/reports/weekly/email", json={"year": YEAR, "squad_id": sa})
    assert r.status_code == 200, r.text
    got = inbox.take()
    assert len(got) == 1
    m = mails["manual"] = _read_mail(got[0], "manual weekly")
    assert MILESTONE_BLOCKED in m["visible"] and "Squad A" in m["subject"]
    # The signals a reader must get without opening the attachment.
    assert OTD_TITLE in m["visible"], "the late commitment is not in the mail"
    assert KPI_NAME in m["visible"], "the KPI under pressure is not in the mail"
    assert KEY_MESSAGE in m["visible"]
    assert OTD_TITLE in m["text"] and KPI_NAME in m["text"], "the text part lacks what the HTML says"
    assert any(n.endswith(".pptx") for n in m["atts"]), m["atts"].keys()
    pptx_name = next(n for n in m["atts"] if n.endswith(".pptx"))
    assert MILESTONE_BLOCKED in _pptx_text(m["atts"][pptx_name])
    assert "https://teamfollowup.example.org" in m["html"], "no link back to the application"

    # 1 bis. The same mail in English says everything in English.
    r = client.post("/api/reports/weekly/email", json={"year": YEAR, "squad_id": sa, "lang": "en"})
    assert r.status_code == 200, r.text
    en = _read_mail(inbox.take()[0], "manual weekly en")
    for label in ("Late commitments", "KPIs to watch", "Milestones to watch", "Key messages", "Open in the app"):
        assert label in en["visible"], f"English mail lacks {label!r}"
    for french in ("Engagements en retard", "Jalons à surveiller", "Ouvrir dans"):
        assert french not in en["visible"], f"French left in the English mail: {french!r}"

    # 2. The scheduled reports: the Direction's whole document, and the tribe's
    # (its leader the tribe's document, each squad leader their squad's).
    login(client, "admin@test")
    now = dt.datetime(YEAR, 6, 15, 9, 0, tzinfo=dt.timezone.utc)
    set_report(db, {"enabled": True, "recipients": ["copil@example.org"], "weekdays": [now.weekday()], "hour": 8})
    set_report(db, {"enabled": True, "weekdays": [now.weekday()], "hour": 8, "copy_leader": True,
                    "leader_mode": "per_squad"}, seeded["t1"])
    db.commit()
    report_mod.send_due_weekly_reports(db, now=now)
    got = inbox.take()
    to_copil = [g for g in got if "copil@example.org" in str(g["To"])]
    # One mail per squad: the tribe leader in To, the squad leader in copy.
    to_leader = [g for g in got if "sl_a@test" in str(g["Cc"]) and "tribe@test" in str(g["To"])]
    assert len(to_copil) == 1, [str(g["Subject"]) for g in got]
    assert to_leader, [str(g["To"]) for g in got]
    for i, g in enumerate(got):
        mails[f"scheduled_{i}"] = _read_mail(g, f"scheduled {i} {g['To']}")
    squad_mail = next(v for k, v in mails.items() if k.startswith("scheduled") and "Squad A" in v["subject"])
    assert MILESTONE_BLOCKED in squad_mail["visible"]

    # 3. The Test button sends a line's mail to the caller.
    r = client.post(f"/api/admin/report-config/test?tribe_id={seeded['t1']}", json={"line": "roles"})
    assert r.status_code == 200 and r.json()["ok"], r.text
    got = inbox.take()
    assert len(got) == 1
    for i, g in enumerate(got):
        _read_mail(g, f"test button {i}")
        assert "(test)" in str(g["Subject"])

    # 4. "Send to the squad leaders" now.
    r = client.post("/api/admin/report-config/send-squad-leaders", json={})
    assert r.status_code == 200, r.text
    for i, g in enumerate(inbox.take()):
        _read_mail(g, f"to leaders {i}")

    # 5. A change notification after the leader edits the reporting.
    login(client, "admin@test")
    set_change_notify(db, {"enabled": True, "recipients": ["copil@example.org"], "events": list(ALL_EVENTS),
                           "attach_pptx": True, "min_interval_minutes": 0})
    db.commit()
    monkeypatch.setattr(dbmod, "SessionLocal", TestingSessionLocal)
    changenotify._run(sa, "roadmap", "SL A", YEAR, "sl_a@test")
    got = inbox.take()
    assert got, "the change notification sent nothing"
    change = _read_mail(got[0], "change notification")
    assert "Squad A" in change["subject"]

    # 7. Access granted: the new user is told, in a mail of its own.
    from app.models import User
    from app.security import hash_password
    db.add(User(email="new.person@example.org", display_name="New Person", role="member",
                status="pending", password_hash=hash_password("pw")))
    db.commit()
    newbie = db.query(User).filter(User.email == "new.person@example.org").one()
    r = client.post(f"/api/access-requests/{newbie.id}/approve", json={"role": "member", "squad_id": sa})
    assert r.status_code == 200, r.text
    got = [g for g in inbox.take() if "new.person@example.org" in str(g["To"])]
    assert got, "the approved person got no mail"
    granted = _read_mail(got[0], "access granted")
    assert "https://teamfollowup.example.org" in granted["html"]

    # 8. The SMTP test mail.
    r = client.post("/api/admin/smtp-config/test", json={"to": "admin@test"})
    assert r.status_code == 200, r.text
    got = inbox.take()
    assert got
    _read_mail(got[0], "smtp test")
