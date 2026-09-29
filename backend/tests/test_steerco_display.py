"""How the steerco slide draws its KPI chart, chosen from the live preview.

With a dozen KPIs, one plot of a dozen curves (on two scales) cannot be read, and
the slide kept five KPI cards and dropped the rest. The platform now chooses how
its chart is drawn, and which KPIs it shows; every KPI keeps its card."""
import io

from app.models import Platform, SteercoEntry
import app.platforms as plat
from app.routers.steerco import I18N, _aggregate, _onepager, _render_pptx
from tests.conftest import login
from tests.test_steerco import _enable, _platform

NAMES = ["Cloud Users", "Landing Zone", "K8aaS", "DBaaS", "Software Factory", "Object Storage",
         "VM Fleet", "Backups", "IAM Roles", "VPN Tunnels", "Load Balancers", "Secrets"]


def _many(db, seeded, n=len(NAMES)) -> int:
    pid = _platform(db, seeded)
    p = db.get(Platform, pid)
    p.template = plat.normalize_template({"kpis": [{"label": x} for x in NAMES[:n]]},
                                         [s.id for s in p.contributors])
    for m in (6, 7):
        db.add(SteercoEntry(platform_id=pid, period=f"2026-{m:02d}", data={"kpis": [
            {"label": x, "value": str((i + 1) * 10 + m)} for i, x in enumerate(NAMES[:n])]}))
    db.commit()
    return pid


def test_auto_draws_curves_for_few_kpis_and_small_charts_beyond(db, seeded):
    few = _aggregate(db, _many(db, seeded, n=4), "2026-07")
    assert few["kpi_chart"]["mode"] == "lines"
    db.query(SteercoEntry).delete(); db.query(Platform).delete(); db.commit()
    many = _aggregate(db, _many(db, seeded), "2026-07")
    assert many["kpi_chart"]["mode"] == "small_multiples"
    # Small charts each have their own scale: nobody is sent to a second axis.
    assert all(s.get("axis") is None for s in many["kpi_chart"]["series"])


def test_the_chart_shows_only_the_kpis_picked_and_keeps_their_colours(db, seeded):
    pid = _many(db, seeded)
    every = {s["name"]: s["color"] for s in _aggregate(db, pid, "2026-07")["kpi_chart"]["series"]}
    rd = _aggregate(db, pid, "2026-07", display={"kpi_chart": "lines", "chart_kpis": ["DBaaS", "Secrets"]})
    shown = rd["kpi_chart"]["series"]
    assert [s["name"] for s in shown] == ["DBaaS", "Secrets"]
    assert all(s["color"] == every[s["name"]] for s in shown)
    # The cards are not the chart: every KPI keeps its card.
    assert len(rd["kpis"]) == len(NAMES)


def test_every_kpi_keeps_its_card_on_the_slide(db, seeded):
    from pptx import Presentation
    rd = _aggregate(db, _many(db, seeded), "2026-07")
    slide = Presentation(io.BytesIO(_render_pptx([{"squad_name": "P", "data": rd}], "2026-07", I18N["en"]))).slides[0]
    text = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
    assert all(x in text for x in NAMES), [x for x in NAMES if x not in text]


def test_each_mode_renders_in_html_and_pptx(db, seeded):
    from pptx import Presentation
    pid = _many(db, seeded)
    for mode, mark in (("small_multiples", "sm-grid"), ("table", "class='kt'"), ("lines", "<svg")):
        rd = _aggregate(db, pid, "2026-07", display={"kpi_chart": mode, "chart_kpis": []})
        assert mark in _onepager("P", "2026-07", rd, I18N["fr"])
        slide = Presentation(io.BytesIO(_render_pptx([{"squad_name": "P", "data": rd}], "2026-07", I18N["fr"]))).slides[0]
        charts = [sh for sh in slide.shapes if sh.has_chart]
        tables = [sh for sh in slide.shapes if sh.has_table]
        if mode == "small_multiples":
            assert len(charts) >= len(NAMES)              # one per KPI (plus incidents, when filled)
        if mode == "table":
            cells = " ".join(c.text for sh in tables for row in sh.table.rows for c in row.cells)
            assert all(x in cells for x in NAMES)         # every KPI has its row


def test_the_rendering_is_saved_by_whoever_manages_platforms(client, db, seeded):
    pid = _many(db, seeded)
    _enable(client)
    login(client, seeded["tribe"])
    r = client.put(f"/api/steerco/platforms/{pid}/display",
                   json={"kpi_chart": "table", "chart_kpis": ["DBaaS", "Unknown KPI"]})
    assert r.status_code == 200, r.text
    assert r.json()["template"]["display"] == {"kpi_chart": "table", "chart_kpis": ["DBaaS"]}
    assert _aggregate(db, pid, "2026-07")["kpi_chart"]["mode"] == "table"
    # Editing the platform's items (a template without "display") keeps the rendering.
    tpl = r.json()["template"]
    del tpl["display"]
    out = client.put(f"/api/steerco/platforms/{pid}", json={"template": tpl}).json()
    assert out["template"]["display"]["kpi_chart"] == "table"
    # An unknown mode falls back to automatic.
    assert client.put(f"/api/steerco/platforms/{pid}/display",
                      json={"kpi_chart": "pie"}).json()["template"]["display"]["kpi_chart"] == "auto"


def test_a_squad_leader_previews_but_does_not_save(client, db, seeded):
    pid = _many(db, seeded)
    _enable(client)
    login(client, seeded["sl_a"])
    assert client.put(f"/api/steerco/platforms/{pid}/display", json={"kpi_chart": "table"}).status_code == 403
    html = client.get(f"/api/steerco/onepager.html?platform_id={pid}&period=2026-07&chart=table"
                      "&chart_kpis=DBaaS&chart_kpis=Secrets").text
    assert "class='kt'" in html and "Secrets" in html


def test_nothing_runs_past_the_first_row_whatever_the_number_of_kpis(db, seeded):
    """Le premier rang (cartes KPI | graphique) ne deborde jamais sur le second,
    quel que soit le nombre de KPI et le rendu choisi: c'est ce qui faisait la
    capture d'origine (legende sur la courbe) et, avec beaucoup de KPI, des cartes
    empilees et des tableaux sous le cadre."""
    from pptx import Presentation
    from pptx.util import Inches
    names = [f"KPI numero {i} avec un nom long" for i in range(30)]
    for n in (8, 20, 30):
        db.query(SteercoEntry).delete(); db.query(Platform).delete(); db.commit()
        pid = _platform(db, seeded)
        p = db.get(Platform, pid)
        p.template = plat.normalize_template({"kpis": [{"label": x} for x in names[:n]]},
                                             [s.id for s in p.contributors])
        for m in range(1, 10):
            db.add(SteercoEntry(platform_id=pid, period=f"2026-{m:02d}", data={"kpis": [
                {"label": x, "value": str(1000000 + i * 12345 + m * 54321)} for i, x in enumerate(names[:n])]}))
        db.commit()
        for mode in ("auto", "lines", "small_multiples", "table"):
            rd = _aggregate(db, pid, "2026-09", display={"kpi_chart": mode, "chart_kpis": []})
            slide = Presentation(io.BytesIO(_render_pptx([{"squad_name": "P", "data": rd}], "2026-09", I18N["fr"]))).slides[0]
            second_row = Inches(3.5)   # top of the SLA / incidents row on the standard slide
            over = [sh.name for sh in slide.shapes
                    if Inches(0.9) < sh.top < second_row and sh.top + sh.height > second_row + Inches(0.02)]
            assert not over, (n, mode, over)
